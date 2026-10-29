"""Sprint 2 product backlog and W02 hour plan (Excel + PDF + markdown)."""

from __future__ import annotations

from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parent
DAYS = [f"Day {i}" for i in range(1, 16)]
NAVY = "1F4E79"
BLUE = "2E75B6"
PALE = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "C6EFCE"

STORIES = [
    {
        "id": "US9a",
        "story": "Edit a logged donation: Given a donor owns a donation that is still queued or matched, when they change the food lines, quantity, expiry, or pickup window, then PATCH /api/donations/:id saves the log, restores the previous food-bank capacity, and runs matching again. After picked_up or delivered, the edit is locked with 409.",
        "tasks": [
            ("Build donor edit form at /donations/[id]/edit", "PRATANIN KAHTEPMUD", 6, [3, 3, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
            ("Implement PATCH /api/donations/:id with capacity restore and rematch", "JIRATHIWAT SUNTIPREEDATHAM", 8, [4, 4, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
            ("QA edit before pickup and 409 after pickup", "SAKKARAJ TUVAJIT", 4, [0, 0, 2, 2, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
        ],
    },
    {
        "id": "US9b",
        "story": "Send one log to many food banks: Given a donor puts one or more foods on a single donation, when splitAcrossBanks is true and the total kilograms exceed one bank, then proposeMatches assigns quantityKg across every in-range bank that still has room, and the sum of quantityKg equals the donation quantity.",
        "tasks": [
            ("Add multi-food lines and the split checkbox on the donor form", "PRATANIN KAHTEPMUD", 6, [0, 0, 3, 3, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
            ("Implement proposeMatches so one donation splits across food banks", "SIRAPASSORN POKASERMSONG", 8, [0, 0, 4, 4, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
            ("Subtract quantityKg from each bank remainingCapacityKg", "JIRATHIWAT SUNTIPREEDATHAM", 6, [0, 0, 0, 0, 3, 3, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
            ("QA one log to many banks (sum of quantityKg, at least two matches)", "SAKKARAJ TUVAJIT", 4, [0, 0, 0, 0, 0, 2, 2, 0, 0, 0, 0, 0, 0, 0, 0]),
        ],
    },
    {
        "id": "US3b",
        "story": "Expiry alerts at 24, 12, and 2 hours: Given a logged donation, when hours until expiry cross 24, 12, or 2, then a notification is stored for staff and for drivers and shown on /notifications.",
        "tasks": [
            ("Build 24/12/2 hour expiry notification job", "JIRATHIWAT SUNTIPREEDATHAM", 6, [0, 0, 0, 0, 3, 3, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
            ("Render expiry alerts on /notifications", "PRATANIN KAHTEPMUD", 4, [0, 0, 0, 0, 0, 0, 2, 2, 0, 0, 0, 0, 0, 0, 0]),
            ("Emit driver-audience notifications for the same thresholds", "SIRAPASSORN POKASERMSONG", 4, [0, 0, 0, 0, 2, 2, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
            ("QA alert thresholds and read state", "SAKKARAJ TUVAJIT", 3, [0, 0, 0, 0, 0, 0, 0, 3, 0, 0, 0, 0, 0, 0, 0]),
        ],
    },
    {
        "id": "US4a",
        "story": "Operations dashboard: Given a signed-in user opens /dashboard, when they filter by donor, food bank id, or status, then GET /api/dashboard shows open, matched, picked up, delivered, and urgent counts.",
        "tasks": [
            ("Build operations board filters", "PRATANIN KAHTEPMUD", 6, [0, 0, 0, 0, 0, 0, 2, 2, 2, 0, 0, 0, 0, 0, 0]),
            ("Implement GET /api/dashboard summary", "JIRATHIWAT SUNTIPREEDATHAM", 5, [0, 0, 0, 0, 0, 0, 0, 2, 3, 0, 0, 0, 0, 0, 0]),
            ("QA dashboard counts and filters", "SAKKARAJ TUVAJIT", 3, [0, 0, 0, 0, 0, 0, 0, 0, 0, 3, 0, 0, 0, 0, 0]),
        ],
    },
    {
        "id": "US4b",
        "story": "Driver activity log: Given a food bank or logistics user opens /drivers, when GET /api/drivers/activity runs, then each driver shows accepted, picked up, and delivered counts plus recent runs.",
        "tasks": [
            ("Implement GET /api/drivers/activity", "SIRAPASSORN POKASERMSONG", 6, [0, 0, 0, 0, 0, 0, 0, 0, 3, 3, 0, 0, 0, 0, 0]),
            ("Build the driver activity page", "PRATANIN KAHTEPMUD", 4, [0, 0, 0, 0, 0, 0, 0, 0, 0, 2, 2, 0, 0, 0, 0]),
            ("QA activity counts after pickup and delivery", "SAKKARAJ TUVAJIT", 3, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 3, 0, 0, 0, 0]),
        ],
    },
    {
        "id": "US5b",
        "story": "Monthly tax summary: Given a donor or logistics user on /receipts, when they export a month, then GET /api/reports/tax returns a CSV or a PDF. This stays Sprint 2 and is not Sprint 1 Done.",
        "tasks": [
            ("Implement monthly tax CSV and PDF", "JIRATHIWAT SUNTIPREEDATHAM", 5, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 2, 3, 0, 0, 0]),
            ("Add Export CSV and Export PDF on the receipts page", "PRATANIN KAHTEPMUD", 3, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 3, 0, 0, 0]),
            ("QA tax export for the selected month", "SAKKARAJ TUVAJIT", 2, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 2, 0, 0]),
        ],
    },
    {
        "id": "US7",
        "story": "Driver GPS ping: Given a driver, when they POST /api/drivers/location with lat and lng, then the ping is stored and shown on the driver activity log within one minute.",
        "tasks": [
            ("Implement POST /api/drivers/location", "SIRAPASSORN POKASERMSONG", 6, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 3, 3, 0, 0]),
            ("Show the latest ping on /drivers", "PRATANIN KAHTEPMUD", 3, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 3, 0, 0]),
            ("QA ping timestamp and coordinates", "SAKKARAJ TUVAJIT", 2, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 2, 0]),
        ],
    },
    {
        "id": "US8",
        "story": "Decay research spike: Given a donation id, when GET /api/decay?donationId= runs, then the response includes hoursRemaining and a heuristic riskScore. This is not a trained model.",
        "tasks": [
            ("Implement GET /api/decay heuristic", "JIRATHIWAT SUNTIPREEDATHAM", 4, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 2, 2, 0]),
            ("QA decay response shape", "SAKKARAJ TUVAJIT", 2, [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 2, 0]),
        ],
    },
]

WINDOWS = [
    ("1", "Mon 2 Nov 2026", "sprint-2-scrum-meeting-1.md", "Kickoff. Day 1 starts after.", 0, "Plans and risks. Every card To-Do. No build hours in the meeting."),
    ("2", "Wed 4 Nov 2026", "sprint-2-scrum-meeting-2.md", "Days 1-2", 14, "US9a form 6/6 In Review. PATCH 8/8 In Review. QA not started. US9b still To-Do."),
    ("3", "Fri 6 Nov 2026", "sprint-2-scrum-meeting-3.md", "Days 3-4", 18, "US9a QA 4/4 In Review. US9b form 6/6 and proposeMatches 8/8 In Review. Capacity and US9b QA still To-Do."),
    ("4", "Mon 9 Nov 2026", "sprint-2-scrum-meeting-4.md", "Days 5-7", 24, "US9a and US9b can move to Done if DoD passes. US3b job 6/6 and driver alerts 4/4. Alerts page 2/4. Dashboard filters 2/6."),
    ("5", "Wed 11 Nov 2026", "sprint-2-scrum-meeting-5.md", "Days 8-9", 17, "US3b QA 3/3 In Review. Dashboard API 5/5. Filters 6/6. Activity API 3/6. US5b, US7, US8 still To-Do."),
    ("6", "Fri 13 Nov 2026", "sprint-2-scrum-meeting-6.md", "Days 10-14 (Day 15 = 0)", 40, "Close. Finish US4a, US4b, US5b, US7, US8. 8 of 8 Done only if DoD passes. Planned 113. Day 15 buffer is 0. No standup after Friday 13."),
]


def day_totals():
    totals = [0] * 15
    effort = 0
    for story in STORIES:
        for _name, _owner, est, hours in story["tasks"]:
            if sum(hours) != est:
                raise SystemExit(f"Hour mismatch: {_name} sum {sum(hours)} != {est}")
            if len(hours) != 15:
                raise SystemExit(f"Need 15 days: {_name}")
            effort += est
            for i, value in enumerate(hours):
                totals[i] += value
    window_hours = [14, 18, 24, 17, 40]
    slices = [totals[0:2], totals[2:4], totals[4:7], totals[7:9], totals[9:14]]
    for expected, actual in zip(window_hours, slices):
        if sum(actual) != expected:
            raise SystemExit(f"Window {expected} != {sum(actual)} from {actual}")
    if totals[14] != 0 or sum(totals) != effort:
        raise SystemExit(f"Totals {totals} effort {effort}")
    owners = {
        "PRATANIN KAHTEPMUD": [0] * 15,
        "JIRATHIWAT SUNTIPREEDATHAM": [0] * 15,
        "SIRAPASSORN POKASERMSONG": [0] * 15,
        "SAKKARAJ TUVAJIT": [0] * 15,
    }
    for story in STORIES:
        for _name, owner, _est, hours in story["tasks"]:
            for index, value in enumerate(hours):
                owners[owner][index] += value
    expected = {
        "PRATANIN KAHTEPMUD": (6, 6, 4, 6, 10, 32),
        "JIRATHIWAT SUNTIPREEDATHAM": (8, 0, 12, 5, 9, 34),
        "SIRAPASSORN POKASERMSONG": (0, 8, 4, 3, 9, 24),
        "SAKKARAJ TUVAJIT": (0, 4, 4, 3, 12, 23),
    }
    slices = [(0, 2), (2, 4), (4, 7), (7, 9), (9, 14)]
    for owner, plan in expected.items():
        actual = [sum(owners[owner][a:b]) for a, b in slices]
        actual.append(sum(owners[owner]))
        if tuple(actual) != plan:
            raise SystemExit(f"{owner} {actual} != {plan}")
    return effort, totals


THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def write_xlsx(effort: int, totals: list[int]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sprint 2 backlog"
    ws.merge_cells("A1:U1")
    title = ws["A1"]
    title.value = "SPRINT 2 BACKLOG"
    title.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.merge_cells("A2:U2")
    ws["A2"].value = (
        "Team QWERTY  |  Sprint 2 of 2  |  2-13 November 2026  |  Ends Friday 13 November  |  "
        "Standups every Monday, Wednesday, and Friday  |  "
        f"Planned {effort} hours  |  Day 15 buffer = 0  |  Status To-Do until that meeting's Definition of Done"
    )
    ws["A2"].font = Font(name="Calibri", size=9, color=WHITE)
    ws["A2"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A2"].alignment = CENTER
    ws.row_dimensions[2].height = 20
    headers = ["ID", "User Story", "Tasks", "Owner", "Status", "Estimated effort", *DAYS]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(3, col, text)
        cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = CENTER
        cell.border = THIN
    ws.row_dimensions[3].height = 22
    ws.auto_filter.ref = "A3:U3"
    ws.freeze_panes = "A4"
    row = 4
    for story in STORIES:
        for index, (name, owner, est, hours) in enumerate(story["tasks"]):
            values = [
                story["id"] if index == 0 else "",
                story["story"] if index == 0 else "",
                name,
                owner,
                "To-Do",
                est,
                *hours,
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col, value)
                cell.border = THIN
                cell.font = Font(name="Calibri", size=9)
                cell.alignment = LEFT if col in (2, 3, 4) else CENTER
                if col == 5:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
            ws.row_dimensions[row].height = 32 if index == 0 else 20
            row += 1
    for col in range(1, 22):
        cell = ws.cell(row, col, "")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        cell.alignment = CENTER
        cell.border = THIN
    ws.cell(row, 5, "TOTAL")
    ws.cell(row, 6, effort)
    for i, value in enumerate(totals, 7):
        ws.cell(row, i, value)
    note = row + 2
    ws.merge_cells(start_row=note, start_column=1, end_row=note + 3, end_column=21)
    ws.cell(
        note,
        1,
        "Planning sheet for Sprint 2. New stories US9a (donor edits a log) and US9b (one log sent to many food banks) "
        "are first. Then US3b, US4a, US4b, US5b, US7, US8. Meetings: Mon 2 Nov kickoff; Wed 4 (Days 1-2, 14h); "
        "Fri 6 (Days 3-4, 18h); Mon 9 (Days 5-7, 24h); Wed 11 (Days 8-9, 17h); "
        "Fri 13 close (Days 10-14, 40h). The sprint ends Friday 13 November. Day 15 = 0. "
        "14+18+24+17+40 = 113. "
        "Do not mark Done before that card's QA and Definition of Done. Sprint 1 stories stay Done and are not rebuilt.",
    )
    ws.cell(note, 1).alignment = LEFT
    ws.cell(note, 1).font = Font(name="Calibri", size=9)
    widths = [8, 42, 42, 28, 12, 12] + [6] * 15
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:3"

    stand = wb.create_sheet("Standup map")
    stand.merge_cells("A1:F1")
    stand["A1"].value = "SPRINT 2  —  STANDUP MAP (Monday, Wednesday, Friday)"
    stand["A1"].font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    stand["A1"].fill = PatternFill("solid", fgColor=NAVY)
    stand["A1"].alignment = CENTER
    stand.row_dimensions[1].height = 24
    headers = ["Meeting", "When", "Script", "Hour-days", "Hours", "Board after the meeting"]
    for col, text in enumerate(headers, 1):
        cell = stand.cell(3, col, text)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        cell.alignment = CENTER
        cell.border = THIN
    for r, item in enumerate(WINDOWS, 4):
        for col, value in enumerate(item, 1):
            cell = stand.cell(r, col, value)
            cell.border = THIN
            cell.alignment = LEFT if col in (3, 6) else CENTER
            cell.font = Font(name="Calibri", size=9)
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
        stand.row_dimensions[r].height = 32
    people = [
        ("Person", "M2 D1-2", "M3 D3-4", "M4 D5-7", "M5 D8-9", "M6 D10-14 close", "Sprint total"),
        ("Pratanin", 6, 6, 4, 6, 10, 32),
        ("Jirathiwat", 8, 0, 12, 5, 9, 34),
        ("Sirapassorn", 0, 8, 4, 3, 9, 24),
        ("Sakkaraj", 0, 4, 4, 3, 12, 23),
        ("TOTAL", 14, 18, 24, 17, 40, 113),
    ]
    start = 13
    stand.merge_cells(start_row=start, start_column=1, end_row=start, end_column=7)
    stand.cell(start, 1, "Hours in each standup window (same cells as the Day 1-15 table). Meeting 1 is the kickoff and logs 0.")
    stand.cell(start, 1).font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for col, text in enumerate(people[0], 1):
        cell = stand.cell(start + 1, col, text)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        cell.alignment = CENTER
        cell.border = THIN
    for r, item in enumerate(people[1:], start + 2):
        for col, value in enumerate(item, 1):
            cell = stand.cell(r, col, value)
            cell.border = THIN
            cell.alignment = CENTER
            cell.font = Font(name="Calibri", size=9, bold=(item[0] == "TOTAL"))
            if item[0] == "TOTAL":
                cell.fill = PatternFill("solid", fgColor=BLUE)
                cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
    for i, width in enumerate([14, 22, 36, 22, 12, 78], 1):
        stand.column_dimensions[get_column_letter(i)].width = width
    stand.column_dimensions["G"].width = 14
    stand.column_dimensions["H"].width = 14
    stand.page_setup.orientation = "landscape"
    stand.page_setup.fitToPage = True
    stand.page_setup.fitToWidth = 1
    stand.page_setup.paperSize = stand.PAPERSIZE_A4
    path = ROOT / "sprint-2-backlog.xlsx"
    wb.save(path)
    return path


class BacklogPdf(FPDF):
    def footer(self) -> None:
        self.set_y(-10)
        self.set_font("Helvetica", size=7)
        self.set_text_color(90, 90, 90)
        self.cell(0, 6, f"W02 Sprint Backlog  |  Sprint 2 plan  |  page {self.page_no()}", align="C")


def write_pdf(effort: int, totals: list[int]) -> Path:
    pdf = BacklogPdf(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=False)
    pdf.add_page()
    left = 6
    usable = pdf.w - 12
    col_w = [11, 46, 52, 32, 14, 16] + [7.05] * 15
    col_w = [w * (usable / sum(col_w)) for w in col_w]
    bottom = pdf.h - 12
    headers = ["ID", "User Story", "Tasks", "Owner", "Status", "Est.", *[str(i) for i in range(1, 16)]]

    def draw_banner():
        pdf.set_xy(left, 7)
        pdf.set_fill_color(31, 78, 121)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(usable, 8, "SPRINT 2 BACKLOG", align="C", fill=True)
        pdf.ln(8)
        pdf.set_x(left)
        pdf.set_fill_color(46, 117, 182)
        pdf.set_font("Helvetica", size=7.5)
        pdf.cell(
            usable,
            5.5,
            f"Team QWERTY  |  2-13 November 2026  |  Ends Friday 13  |  15 hour-days  |  Mon / Wed / Fri  |  Planned {effort}h  |  Day 15 = 0",
            align="C",
            fill=True,
        )
        pdf.ln(6)

    def measure(values, fonts, line_h=3.1) -> float:
        max_h = line_h + 1
        for i, text in enumerate(values):
            pdf.set_font(*fonts[i])
            lines = pdf.multi_cell(col_w[i] - 0.8, line_h, str(text), dry_run=True, output="LINES")
            max_h = max(max_h, line_h * max(len(lines), 1) + 1.0)
        return max_h

    def paint_row(values, fills, fonts, text_colors, min_h=7.0):
        height = max(min_h, measure(values, fonts))
        if pdf.get_y() + height > bottom:
            pdf.add_page()
            draw_banner()
            paint_row(headers, [(31, 78, 121)] * 21, [("Helvetica", "B", 5.5)] * 21, [(255, 255, 255)] * 21, 8)
            height = max(min_h, measure(values, fonts))
        y0 = pdf.get_y()
        for i, text in enumerate(values):
            x = left + sum(col_w[:i])
            pdf.set_fill_color(*fills[i])
            pdf.set_draw_color(0, 0, 0)
            pdf.rect(x, y0, col_w[i], height, style="DF")
            pdf.set_xy(x + 0.35, y0 + 0.35)
            pdf.set_text_color(*text_colors[i])
            pdf.set_font(*fonts[i])
            align = "C" if i not in (1, 2, 3) else "L"
            pdf.multi_cell(col_w[i] - 0.7, 3.1, str(text), align=align)
        pdf.set_y(y0 + height)

    draw_banner()
    paint_row(headers, [(31, 78, 121)] * 21, [("Helvetica", "B", 5.5)] * 21, [(255, 255, 255)] * 21, 8)
    alt = [(255, 255, 255), (242, 242, 242)]
    black = (20, 20, 20)
    for story_index, story in enumerate(STORIES):
        bg = alt[story_index % 2]
        for i, (name, owner, est, hours) in enumerate(story["tasks"]):
            values = [story["id"] if i == 0 else "", story["story"] if i == 0 else "", name, owner, "To-Do", est, *hours]
            fills = [bg] * 21
            fills[4] = (255, 242, 204)
            fonts = [("Helvetica", "", 5.6)] * 21
            fonts[4] = ("Helvetica", "B", 6)
            colors = [black] * 21
            colors[4] = (102, 60, 0)
            paint_row(values, fills, fonts, colors, 8 if i == 0 else 6.2)
    paint_row(
        ["", "", "", "", "TOTAL", effort, *totals],
        [(46, 117, 182)] * 21,
        [("Helvetica", "B", 6.5)] * 21,
        [(255, 255, 255)] * 21,
        7,
    )
    path = ROOT / "sprint-2-backlog.pdf"
    pdf.output(path)
    return path


def write_md(effort: int, totals: list[int]) -> Path:
    lines = [
        "# Sprint 2 product backlog and sprint backlog",
        "",
        "Team QWERTY. Sprint 2 of 2. **Monday 2 November 2026 through Friday 13 November 2026**. The sprint ends on Friday 13. Daily Scrum is every **Monday, Wednesday, and Friday** in those two weeks. The W02 sheet still has **15 hour-days**. Day 15's build cell is 0 and is reported at the Friday 13 close. There is no standup after Friday 13. Planned hours **113**. Status on this sheet is **To-Do** until that card's Definition of Done passes at its meeting.",
        "",
        "Sprint 1 stays Done (US1a, US1b, US6, US2a, US2b, US5a, US3a). This sprint adds two stories the donor asked for, then the remaining product backlog.",
        "",
        "Printable copies: `sprint-2-backlog.xlsx` and `sprint-2-backlog.pdf`.",
        "",
        "## Stories",
        "",
        "| ID | User story |",
        "|---|---|",
    ]
    for story in STORIES:
        lines.append(f"| {story['id']} | {story['story']} |")
    lines += [
        "",
        "## Standup windows",
        "",
        "Three Daily Scrums a week. Speak only the hours in that window. Do not call a later card Done.",
        "",
        "| Meeting | When | Script | Hour-days | Hours | Board after the meeting |",
        "|---|---|---|---|---:|---|",
    ]
    for item in WINDOWS:
        lines.append(f"| {item[0]} | {item[1]} | `{item[2]}` | {item[3]} | {item[4]} | {item[5]} |")
    lines += [
        "",
        "Window check: 0 + 14 + 18 + 24 + 17 + 40 = **113**. Day totals "
        + " + ".join(str(n) for n in totals)
        + f" = **{effort}**.",
        "",
        "## Hours by person and meeting",
        "",
        "| Person | M2 Days 1-2 | M3 Days 3-4 | M4 Days 5-7 | M5 Days 8-9 | M6 Days 10-14 (close) | Sprint total |",
        "|---|---:|---:|---:|---:|---:|---:|",
        "| Pratanin | 6 | 6 | 4 | 6 | 10 | 32 |",
        "| Jirathiwat | 8 | 0 | 12 | 5 | 9 | 34 |",
        "| Sirapassorn | 0 | 8 | 4 | 3 | 9 | 24 |",
        "| Sakkaraj | 0 | 4 | 4 | 3 | 12 | 23 |",
        "| **TOTAL** | **14** | **18** | **24** | **17** | **40** | **113** |",
        "",
        "## SPRINT 2 BACKLOG",
        "",
        "| ID | User Story | Tasks | Owner | Status | Estimated effort | "
        + " | ".join(DAYS)
        + " |",
        "|---|---|---|---|---|---:|"
        + "|".join(["---:"] * 15)
        + "|",
    ]
    for story in STORIES:
        for index, (name, owner, est, hours) in enumerate(story["tasks"]):
            sid = story["id"] if index == 0 else ""
            text = story["story"] if index == 0 else ""
            cells = " | ".join(str(n) for n in hours)
            lines.append(f"| {sid} | {text} | {name} | {owner} | To-Do | {est} | {cells} |")
    lines.append(
        "| | | | | **TOTAL** | **"
        + str(effort)
        + "** | "
        + " | ".join(f"**{n}**" for n in totals)
        + " |"
    )
    lines += [
        "",
        "## Definition of Done",
        "",
        "A card is Done only when acceptance criteria pass, the code is reviewed and on `main`, tests pass, validation exists, the time targets still hold (queue under 1 minute, match under 2 minutes), QA has no critical bugs, and this board's hours are updated.",
        "",
        "US1a fields stay stable: `donationId`, `foodType`, `quantity` (kg), `expiryDateTime`, `pickupWindow`, `donorId`, `timezone: Asia/Bangkok`. A multi-food log still exposes those fields. `quantity` is the sum of `lines`. `foodType` is the line list. `expiryDateTime` is the soonest expiry.",
        "",
    ]
    path = ROOT / "sprint-2-backlog.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main() -> None:
    effort, totals = day_totals()
    print(write_md(effort, totals))
    print(write_xlsx(effort, totals))
    print(write_pdf(effort, totals))


if __name__ == "__main__":
    main()
